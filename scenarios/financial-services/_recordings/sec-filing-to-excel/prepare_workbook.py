#!/usr/bin/env python3
"""Developer-only preparation of the blank recording workbook for the SEC-filing-to-Excel pilot.

This script is a local authoring helper. It is not part of any Creator runtime, it is not
bundled with a plugin, it installs nothing, it touches no COM/native automation, it changes no
system state, and it performs no network access. It only writes one .xlsx file to the path you
pass on the command line.

What it produces is an EMPTY, formatted template. Every monetary cell and every calculation cell
is deliberately left blank so that the recorded session shows an operator entering the values
and formulas in the real Excel UI. The template therefore contains no formulas at all, which also
means it carries no recalculation dependency: there is nothing to recalculate and no cached
formula value to fake.

Usage:
    python prepare_workbook.py --output initial.xlsx
    python prepare_workbook.py --output initial.xlsx --force

Requires openpyxl (3.1.5 verified). openpyxl is a developer-only dependency of THIS script; the
stdlib validator and the stdlib unit tests in this folder do not import it.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
import zipfile

try:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:  # pragma: no cover - developer environment guard
    sys.stderr.write(
        "prepare_workbook.py needs openpyxl (developer-only). Install it in your own environment "
        "or run this script where openpyxl 3.1.5 is already present.\n"
    )
    raise SystemExit(2)

TEMPLATE_TOKEN = "KO-SOI-RECORDING-TEMPLATE-v1"

FONT_NAME = "Arial"
SIZE_TITLE = 16
SIZE_SUBTITLE = 12
SIZE_DATA = 14
SIZE_NOTE = 11

BLUE = "FF0000FF"
BLACK = "FF000000"
GREEN = "FF008000"
RED = "FFFF0000"
GREY = "FF595959"
YELLOW = "FFFFFF00"

MONEY_FORMAT = '#,##0;(#,##0);"-"'
PERCENT_FORMAT = '0.0%;(0.0%);"-"'

THIN = Side(style="thin", color="FFBFBFBF")
HEADER_RULE = Border(bottom=Side(style="medium", color="FF000000"))
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Row 5..10 labels, in the fixed order agreed for the recording.
STATEMENT_ROWS = [
    (5, "Net Operating Revenues"),
    (6, "Cost of goods sold"),
    (7, "Gross Profit"),
    (8, "Selling, general and administrative expenses"),
    (9, "Other operating charges"),
    (10, "Operating Income"),
]

# Derived block: (row, label, intended B formula, intended C formula, number format).
# These formulas are DOCUMENTATION ONLY. They are not written into the template.
DERIVED_ROWS = [
    (12, "Recomputed gross profit", "=B5-B6", "=C5-C6", MONEY_FORMAT),
    (13, "Gross profit difference", "=B12-B7", "=C12-C7", MONEY_FORMAT),
    (14, "Recomputed operating income", "=B12-B8-B9", "=C12-C8-C9", MONEY_FORMAT),
    (15, "Operating income difference", "=B14-B10", "=C14-C10", MONEY_FORMAT),
    (17, "Gross margin", "=B12/B5", "=C12/C5", PERCENT_FORMAT),
    (18, "Operating margin", "=B14/B5", "=C14/C5", PERCENT_FORMAT),
]

CONTROL_FORMULA = (
    '=IF(COUNT(B5:C10)<12,"Pending inputs",'
    'IF(AND(B13=0,C13=0,B15=0,C15=0),"Ties to filing","Review difference"))'
)

SOURCE_CITATION = (
    "Source: The Coca-Cola Company FY2024 Form 10-K, Item 8, Consolidated Statements of Income, "
    "printed page 62, accession 0000021344-25-000011, "
    "https://www.sec.gov/Archives/edgar/data/21344/000002134425000011/ko-20241231.htm"
)

SOURCE_FIELDS = [
    ("Company", "The Coca-Cola Company and Subsidiaries"),
    ("Ticker", "KO"),
    ("CIK", "0000021344"),
    ("Form type", "Form 10-K (annual report)"),
    ("Fiscal year ended", "2024-12-31"),
    ("Accession number", "0000021344-25-000011"),
    (
        "Primary document URL",
        "https://www.sec.gov/Archives/edgar/data/21344/000002134425000011/ko-20241231.htm",
    ),
    ("Statement", "Item 8, Consolidated Statements of Income"),
    ("Printed page", "62"),
    ("In-document anchor observed", "#i9dbdca2a3d754daa96efadd0a7b755b5_124"),
    ("Heading immediately above table", "THE COCA-COLA COMPANY AND SUBSIDIARIES"),
    ("Statement caption", "CONSOLIDATED STATEMENTS OF INCOME"),
    ("Units caption", "(In millions except per share data)"),
    ("Column header in source", "Year Ended December 31, | 2024 | 2023 | 2022"),
    ("Columns transferred to Statement", "2024 and 2023 (both from this same filing)"),
    ("Column not transferred", "2022 (present in source table, out of scope for this pilot)"),
    ("Filing date", "2025-02-20 (read from the EDGAR filing index)"),
    ("Period of report", "2024-12-31"),
    (
        "Filing index URL",
        "https://www.sec.gov/Archives/edgar/data/21344/000002134425000011/"
        "0000021344-25-000011-index.html",
    ),
    ("Date accessed", "2026-09-16"),
    ("Retrieval method", "Original SEC filing read in a browser with UI-assisted observation"),
    (
        "Table deliberately EXCLUDED",
        "The separate equity-method investee summary table later in the same 10-K repeats similar "
        "revenue and gross profit captions with different amounts. It is not the consolidated "
        "company statement and must not be used.",
    ),
    (
        "Amounts",
        "Intentionally not stored on this sheet. Amounts are typed into 'Statement' during the "
        "recorded session and are cited back to the reference above.",
    ),
]

GUIDE_STEPS = [
    "Confirm the source table first: company heading, 'CONSOLIDATED STATEMENTS OF INCOME', the "
    "'(In millions except per share data)' units caption, and the 'Year Ended December 31,' "
    "column header.",
    "Reject the equity-method investee summary table. It reuses similar captions with different "
    "amounts and is not the consolidated company statement.",
    "Paste the original statement table into 'Raw paste' starting at A5. Leave rows 1-4 as the "
    "sheet's own title and note.",
    "Type the six 2024 amounts into B5:B10 and the six 2023 amounts into C5:C10. Blue text on a "
    "yellow fill marks values transcribed from the filing.",
    "Type the change formulas in D5:D10 and the guarded percent-change formulas in E5:E10.",
    "Type the recomputation block in B12:C15, the margins in B17:C18, and the control in B20.",
    "Read the control cell and the two difference rows. Report what they actually say. Do not "
    "edit a reported amount to force a tie.",
    "Record any question on the 'Review' sheet with its evidence and disposition.",
    "Save the working draft locally with Ctrl+S. A human reviews the result; nothing is posted, "
    "paid, traded, or issued as advice.",
]

GUIDE_LEGEND = [
    ("Blue text on yellow fill", "Typed input taken from the filing", BLUE, YELLOW),
    ("Black text", "Formula or calculation", BLACK, None),
    ("Green text", "Link to another sheet in this workbook", GREEN, None),
    ("Red text", "Link to another file (not used in this pilot)", RED, None),
]


def text_cell(worksheet, coord, value):
    """Write a literal string, even when it starts with '=' so Excel keeps it as text."""
    cell = worksheet[coord]
    cell.value = value
    cell.data_type = "s"
    return cell


def style(cell, *, size=SIZE_DATA, bold=False, color=BLACK, italic=False,
          fill=None, horizontal=None, vertical="center", wrap=False, number_format=None,
          border=None):
    cell.font = Font(name=FONT_NAME, size=size, bold=bold, italic=italic, color=color)
    cell.alignment = Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)
    if fill is not None:
        cell.fill = PatternFill(fill_type="solid", start_color=fill, end_color=fill)
    if number_format is not None:
        cell.number_format = number_format
    if border is not None:
        cell.border = border
    return cell


def sheet_title(worksheet, title, note, last_column="E"):
    worksheet.merge_cells(f"A1:{last_column}1")
    style(worksheet["A1"], size=SIZE_TITLE, bold=True, horizontal="left")
    worksheet["A1"] = title
    worksheet.merge_cells(f"A2:{last_column}2")
    style(worksheet["A2"], size=SIZE_SUBTITLE, color=GREY, horizontal="left", wrap=True)
    worksheet["A2"] = note
    worksheet.row_dimensions[1].height = 24
    worksheet.row_dimensions[2].height = 32


def build_statement(worksheet):
    sheet_title(
        worksheet,
        "The Coca-Cola Company and Subsidiaries (KO) - Consolidated Statements of Income",
        "USD in millions, except per share data. Fiscal years ended December 31. Amounts are "
        "transferred from the FY2024 Form 10-K, Item 8; see the 'Sources' sheet for the full "
        "citation. Amounts and formulas are entered during the recorded session.",
    )

    headers = [
        ("A4", "Metric"),
        ("B4", "2024 ($mm)"),
        ("C4", "2023 ($mm)"),
        ("D4", "Change ($mm)"),
        ("E4", "Change (%)"),
    ]
    for coord, label in headers:
        cell = worksheet[coord]
        cell.value = label
        align = "left" if coord == "A4" else "right"
        style(cell, bold=True, horizontal=align, wrap=True, border=HEADER_RULE)
    worksheet.row_dimensions[4].height = 34

    for row, label in STATEMENT_ROWS:
        style(worksheet.cell(row=row, column=1, value=label), horizontal="left", wrap=True)
        for column in (2, 3):
            style(
                worksheet.cell(row=row, column=column),
                color=BLUE,
                fill=YELLOW,
                horizontal="right",
                number_format=MONEY_FORMAT,
                border=BOX,
            )
        style(worksheet.cell(row=row, column=4), horizontal="right",
              number_format=MONEY_FORMAT, border=BOX)
        style(worksheet.cell(row=row, column=5), horizontal="right",
              number_format=PERCENT_FORMAT, border=BOX)
        worksheet.row_dimensions[row].height = 22

    for row, label, _b, _c, number_format in DERIVED_ROWS:
        style(worksheet.cell(row=row, column=1, value=label), horizontal="left", wrap=True)
        for column in (2, 3):
            style(worksheet.cell(row=row, column=column), horizontal="right",
                  number_format=number_format, border=BOX)
        worksheet.row_dimensions[row].height = 22

    style(worksheet.cell(row=20, column=1, value="Control"), bold=True, horizontal="left")
    style(worksheet.cell(row=20, column=2), horizontal="left", border=BOX)
    worksheet.row_dimensions[20].height = 22

    style(
        worksheet.cell(
            row=22,
            column=1,
            value="Status: Pending inputs. Workbook prepared blank; no amounts or formulas "
                  "entered yet.",
        ),
        size=SIZE_NOTE,
        italic=True,
        color=GREY,
        horizontal="left",
    )
    style(worksheet.cell(row=23, column=1, value=SOURCE_CITATION),
          size=SIZE_NOTE, color=GREY, horizontal="left", wrap=True)
    worksheet.row_dimensions[23].height = 30

    worksheet["A4"].comment = Comment(
        "Rows 5-10 mirror six captions of the consolidated statement, in statement order.\n"
        + SOURCE_CITATION,
        "Pilot preparation script",
        height=170,
        width=460,
    )
    worksheet["B4"].comment = Comment(
        "Type the 2024 column of the consolidated statement here (USD millions).\n"
        "Do not use the equity-method investee summary table.",
        "Pilot preparation script",
        height=130,
        width=420,
    )
    worksheet["C4"].comment = Comment(
        "Type the 2023 column from the SAME filing here (USD millions).\n"
        "Do not substitute a prior-year filing.",
        "Pilot preparation script",
        height=130,
        width=420,
    )
    worksheet["A20"].comment = Comment(
        'Control reads "Pending inputs" until all twelve amounts are present, then reports '
        'whether the recomputation ties to the reported subtotals.',
        "Pilot preparation script",
        height=130,
        width=420,
    )

    for column, width in (("A", 52), ("B", 18), ("C", 18), ("D", 18), ("E", 16)):
        worksheet.column_dimensions[column].width = width
    worksheet.freeze_panes = "B5"
    worksheet.sheet_view.showGridLines = False


def build_raw_paste(worksheet):
    sheet_title(
        worksheet,
        "Raw paste - original statement table as copied from the filing",
        "Paste the copied source table starting at A5 and leave it unedited. The source HTML "
        "contains standalone '$' cells, blank spacer cells, non-breaking spaces after "
        "comma-formatted numbers, and a 2022 column; that is expected and is cleaned on the "
        "'Statement' sheet, not here.",
        last_column="H",
    )
    style(worksheet.cell(row=4, column=1, value="Paste below this line"),
          size=SIZE_NOTE, italic=True, color=GREY, horizontal="left")
    for column in ("A", "B", "C", "D", "E", "F", "G", "H"):
        worksheet.column_dimensions[column].width = 26 if column == "A" else 16
    worksheet.freeze_panes = "A5"


def build_sources(worksheet):
    sheet_title(
        worksheet,
        "Sources - filing identity and citation",
        "Source-checked reference metadata recorded before the session. No reported amounts are "
        "stored on this sheet.",
        last_column="B",
    )
    style(worksheet.cell(row=4, column=1, value="Field"), bold=True,
          horizontal="left", border=HEADER_RULE)
    style(worksheet.cell(row=4, column=2, value="Value"), bold=True,
          horizontal="left", border=HEADER_RULE)
    row = 5
    for label, value in SOURCE_FIELDS:
        style(worksheet.cell(row=row, column=1, value=label), bold=True, horizontal="left",
              vertical="top", wrap=True)
        style(worksheet.cell(row=row, column=2, value=value), horizontal="left",
              vertical="top", wrap=True)
        worksheet.row_dimensions[row].height = 44 if len(value) > 70 else 22
        row += 1
    worksheet.column_dimensions["A"].width = 38
    worksheet.column_dimensions["B"].width = 86
    worksheet.freeze_panes = "A5"


def build_review(worksheet):
    sheet_title(
        worksheet,
        "Review - questions raised during the session",
        "Fill in only what is actually observed. Do not pre-seed an issue, and do not invent a "
        "discrepancy with the issuer's reported figures.",
        last_column="D",
    )
    for index, label in enumerate(("Issue", "Evidence", "Disposition", "Reviewer status"), start=1):
        style(worksheet.cell(row=4, column=index, value=label), bold=True,
              horizontal="left", border=HEADER_RULE)
    for row in range(5, 13):
        for column in range(1, 5):
            style(worksheet.cell(row=row, column=column), horizontal="left",
                  vertical="top", wrap=True, border=BOX)
        worksheet.row_dimensions[row].height = 30
    for column, width in (("A", 40), ("B", 52), ("C", 40), ("D", 22)):
        worksheet.column_dimensions[column].width = width
    worksheet.freeze_panes = "A5"


def build_guide(worksheet):
    sheet_title(
        worksheet,
        "Guide - analyst steps and formula references",
        "Working notes for the person doing the task. Formula text below is stored as text, not "
        "as live formulas, and no resulting amounts are pre-filled.",
        last_column="C",
    )
    style(worksheet.cell(row=4, column=1, value="Step"), bold=True,
          horizontal="left", border=HEADER_RULE)
    style(worksheet.cell(row=4, column=2, value="Action"), bold=True,
          horizontal="left", border=HEADER_RULE)
    row = 5
    for index, step in enumerate(GUIDE_STEPS, start=1):
        style(worksheet.cell(row=row, column=1, value=index), horizontal="left", vertical="top")
        style(worksheet.cell(row=row, column=2, value=step), horizontal="left",
              vertical="top", wrap=True)
        worksheet.row_dimensions[row].height = 42
        row += 1

    row += 1
    style(worksheet.cell(row=row, column=1, value="Cell"), bold=True,
          horizontal="left", border=HEADER_RULE)
    style(worksheet.cell(row=row, column=2, value="Formula to type on 'Statement'"), bold=True,
          horizontal="left", border=HEADER_RULE)
    style(worksheet.cell(row=row, column=3, value="Meaning"), bold=True,
          horizontal="left", border=HEADER_RULE)
    row += 1

    formula_rows = []
    for statement_row, _label in STATEMENT_ROWS:
        formula_rows.append(
            (f"D{statement_row}", f"=B{statement_row}-C{statement_row}", "Year-over-year change")
        )
    for statement_row, _label in STATEMENT_ROWS:
        formula_rows.append((
            f"E{statement_row}",
            f'=IF(C{statement_row}=0,"Not meaningful",D{statement_row}/C{statement_row})',
            "Percent change, guarded only against a true zero prior-year base",
        ))
    for derived_row, label, b_formula, c_formula, _fmt in DERIVED_ROWS:
        formula_rows.append((f"B{derived_row}", b_formula, label + " (2024)"))
        formula_rows.append((f"C{derived_row}", c_formula, label + " (2023)"))
    formula_rows.append(("B20", CONTROL_FORMULA, "Control status for the whole transfer"))

    for coord, formula, meaning in formula_rows:
        style(worksheet.cell(row=row, column=1, value=coord), horizontal="left", vertical="top")
        style(text_cell(worksheet, f"B{row}", formula), horizontal="left", vertical="top",
              wrap=True)
        style(worksheet.cell(row=row, column=3, value=meaning), horizontal="left",
              vertical="top", wrap=True)
        worksheet.row_dimensions[row].height = 30
        row += 1

    row += 1
    style(worksheet.cell(row=row, column=1, value="Colour"), bold=True,
          horizontal="left", border=HEADER_RULE)
    style(worksheet.cell(row=row, column=2, value="Convention"), bold=True,
          horizontal="left", border=HEADER_RULE)
    row += 1
    for sample, meaning, color, fill in GUIDE_LEGEND:
        style(worksheet.cell(row=row, column=1, value=sample), color=color, fill=fill,
              horizontal="left", vertical="top", wrap=True)
        style(worksheet.cell(row=row, column=2, value=meaning), horizontal="left",
              vertical="top", wrap=True)
        worksheet.row_dimensions[row].height = 26
        row += 1

    row += 1
    style(worksheet.cell(row=row, column=1, value=SOURCE_CITATION),
          size=SIZE_NOTE, color=GREY, horizontal="left", wrap=True)
    worksheet.row_dimensions[row].height = 34

    for column, width in (("A", 34), ("B", 74), ("C", 46)):
        worksheet.column_dimensions[column].width = width
    worksheet.freeze_panes = "A5"


def build_workbook():
    workbook = Workbook()
    statement = workbook.active
    statement.title = "Statement"
    build_statement(statement)
    build_raw_paste(workbook.create_sheet("Raw paste"))
    build_sources(workbook.create_sheet("Sources"))
    build_review(workbook.create_sheet("Review"))
    build_guide(workbook.create_sheet("Guide"))
    workbook.active = 0

    fixed = datetime(2026, 1, 1, tzinfo=timezone.utc).replace(tzinfo=None)
    properties = workbook.properties
    properties.creator = "Financial reporting pilot - workbook preparation script"
    properties.lastModifiedBy = "Financial reporting pilot - workbook preparation script"
    properties.title = "KO consolidated statements of income - recording template"
    properties.subject = "Blank template for a recorded SEC-filing-to-Excel transfer"
    properties.category = "Template"
    properties.keywords = TEMPLATE_TOKEN
    properties.description = (
        f"{TEMPLATE_TOKEN}. Blank preparation template. Amounts and formulas are entered in the "
        "Excel user interface during the recorded session; nothing is pre-computed here."
    )
    properties.created = fixed
    properties.modified = fixed
    return workbook


def existing_is_same_template(path: Path) -> bool:
    """True only when an existing file is a previous build of this same template."""
    try:
        with zipfile.ZipFile(path) as archive:
            core = archive.read("docProps/core.xml").decode("utf-8", "replace")
    except (OSError, KeyError, zipfile.BadZipFile):
        return False
    try:
        root = ET.fromstring(core)
    except ET.ParseError:
        return False
    return any(
        element.text and TEMPLATE_TOKEN in element.text
        for element in root.iter()
        if element.text
    )


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Developer-only: write the blank, formatted recording template. No network, no COM, "
            "no installs, no system changes."
        )
    )
    parser.add_argument("--output", required=True,
                        help="Explicit path of the .xlsx file to write.")
    parser.add_argument("--force", action="store_true",
                        help="Overwrite an existing file that is not a build of this template.")
    arguments = parser.parse_args(argv)

    output = Path(arguments.output)
    if output.suffix.lower() != ".xlsx":
        sys.stderr.write(f"Refusing to write a non-.xlsx output path: {output}\n")
        return 2
    if output.exists():
        if not output.is_file():
            sys.stderr.write(f"Refusing to overwrite a non-file path: {output}\n")
            return 2
        if not arguments.force and not existing_is_same_template(output):
            sys.stderr.write(
                f"Refusing to overwrite an existing file that is not a build of "
                f"{TEMPLATE_TOKEN}: {output}\nRe-run with --force if that is really intended.\n"
            )
            return 2
    if output.parent and not output.parent.exists():
        sys.stderr.write(f"Output directory does not exist: {output.parent}\n")
        return 2

    build_workbook().save(output)
    print(f"Wrote blank recording template: {output}")
    print("Statement!B5:C10, D5:E10, B12:C15, B17:C18 and B20 are intentionally empty.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
